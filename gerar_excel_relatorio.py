"""
Gera Excel de monitorias no mesmo padrão dos relatórios de referência.
Abas: Resumo | Por Ligação | Por Gestor | Falhas Comuns

Uso:
    python gerar_excel_relatorio.py                  -> último output gerado
    python gerar_excel_relatorio.py 2026-05-26       -> data específica
"""

import sys
import json
import os
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR   = r"G:\Meu Drive\MONITORIAS RAGAZ 05.26"
OUTPUT_DIR = os.path.join(BASE_DIR, "output")

# Cores
COR_HEADER      = "1F4E78"
COR_SUBHEADER   = "2E75B6"
COR_LINHA_PAR   = "EBF3FB"
COR_VERDE       = "E2EFDA"
COR_VERMELHO    = "FCE4D6"
COR_AMARELO     = "FFF2CC"
COR_CINZA       = "F2F2F2"

def _style_header(cell, cor=COR_HEADER, fonte_cor="FFFFFF", tamanho=11, bold=True):
    cell.font = Font(bold=bold, color=fonte_cor, size=tamanho, name="Arial")
    cell.fill = PatternFill(start_color=cor, end_color=cor, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def _style_titulo(cell, texto):
    cell.value = texto
    cell.font = Font(bold=True, size=13, name="Arial", color=COR_HEADER)
    cell.alignment = Alignment(horizontal="left", vertical="center")

def _bordas(cell):
    borda = Side(style="thin", color="BFBFBF")
    cell.border = Border(left=borda, right=borda, top=borda, bottom=borda)

def _cor_efetividade(valor):
    if valor == "SIM":
        return COR_VERDE
    elif valor == "NAO":
        return COR_VERMELHO
    elif valor == "PARCIAL":
        return COR_AMARELO
    return None

def _cor_coaching(valor):
    if valor == "ALTA":
        return COR_VERMELHO
    elif valor == "MEDIA":
        return COR_AMARELO
    elif valor == "BAIXA":
        return COR_VERDE
    return None


# ── ABA RESUMO ─────────────────────────────────────────────────────────────────

def aba_resumo(wb, relatorio):
    ws = wb.create_sheet("Resumo")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 35
    ws.column_dimensions['B'].width = 20
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 20

    area     = relatorio.get("area", "")
    periodo  = relatorio.get("periodo", "")
    superv   = relatorio.get("supervisor", "")
    consol   = relatorio.get("consolidado", {})
    gerado   = relatorio.get("gerado_em", "")[:10]

    # Título
    ws.merge_cells("A1:D1")
    _style_titulo(ws["A1"], f"MONITORIA DE QUALIDADE — {area.upper()}")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:D2")
    ws["A2"].value = f"Período: {periodo}   |   Supervisor: {superv}   |   Gerado em: {gerado}"
    ws["A2"].font = Font(size=10, color="595959", name="Arial")
    ws["A2"].alignment = Alignment(horizontal="left")
    ws.row_dimensions[2].height = 18

    # Espaço
    ws.row_dimensions[3].height = 8

    # INDICADORES GERAIS
    ws.merge_cells("A4:D4")
    ws["A4"].value = "INDICADORES GERAIS"
    ws["A4"].font = Font(bold=True, size=11, color="FFFFFF", name="Arial")
    ws["A4"].fill = PatternFill(start_color=COR_SUBHEADER, end_color=COR_SUBHEADER, fill_type="solid")
    ws["A4"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws.row_dimensions[4].height = 20

    indicadores = [
        ("Total de ligações (Excel)",    consol.get("total_ligacoes", 0)),
        ("Analisadas pelo Claude",        consol.get("analisadas", 0)),
        ("Cobertura",                     f"{consol.get('cobertura_pct', 0):.1f}%"),
        ("Custo da análise (USD)",        f"${consol.get('custo_analise_usd', 0):.4f}"),
    ]

    # Campos específicos por área
    area_key = relatorio.get("area_key", "")
    if area_key == "SUCESSO_CORRETOR":
        indicadores += [
            ("Ligações efetivas",         consol.get("efetivas", 0)),
            ("Taxa de efetividade",        f"{consol.get('efetividade_pct', 0):.1f}%"),
            ("Coaching prioridade ALTA",   consol.get("coaching_alta", 0)),
        ]
    elif area_key == "ATENDIMENTO":
        indicadores += [
            ("Nota média",                f"{consol.get('nota_media', 0):.1f}"),
            ("Acima de 70",               consol.get("acima_70", 0)),
            ("Abaixo de 50",              consol.get("abaixo_50", 0)),
            ("Coaching prioridade ALTA",  consol.get("coaching_alta", 0)),
        ]
    elif area_key == "COMERCIAL_CADASTRO":
        indicadores += [
            ("Convertidas",               consol.get("convertidas", 0)),
            ("Taxa de conversão",         f"{consol.get('conversao_pct', 0):.1f}%"),
            ("Coaching prioridade ALTA",  consol.get("coaching_alta", 0)),
        ]

    row = 5
    for i, (label, valor) in enumerate(indicadores):
        ws[f"A{row}"] = label
        ws[f"B{row}"] = valor
        ws[f"A{row}"].font = Font(name="Arial", size=10)
        ws[f"B{row}"].font = Font(bold=True, name="Arial", size=10)
        ws[f"B{row}"].alignment = Alignment(horizontal="center")
        if i % 2 == 0:
            ws[f"A{row}"].fill = PatternFill(start_color=COR_CINZA, end_color=COR_CINZA, fill_type="solid")
            ws[f"B{row}"].fill = PatternFill(start_color=COR_CINZA, end_color=COR_CINZA, fill_type="solid")
        for col in ["A", "B"]:
            _bordas(ws[f"{col}{row}"])
        row += 1

    return ws


# ── ABA POR LIGAÇÃO ────────────────────────────────────────────────────────────

def aba_por_ligacao(wb, relatorio):
    ws = wb.create_sheet("Por Ligação")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    area_key  = relatorio.get("area_key", "")
    ligacoes  = relatorio.get("ligacoes", [])

    # Colunas por área
    if area_key == "SUCESSO_CORRETOR":
        headers = ["ID", "Agente", "Data", "Duração (s)", "Efetivo", "Coaching",
                   "Tipo Pendência", "Padrão Comercial", "Impacto Receita",
                   "SPIN S", "SPIN P", "SPIN I", "SPIN N",
                   "Conformidade", "Resumo"]
        widths  = [32, 22, 12, 12, 10, 10, 20, 15, 15, 8, 8, 8, 8, 12, 60]

    elif area_key == "ATENDIMENTO":
        headers = ["ID", "Agente", "Data", "Duração (s)", "Nota", "Coaching",
                   "Saudação", "Atenção", "Solução Efetiva", "Clareza",
                   "Transferência", "Tempo Espera", "Vocabulário", "Cordialidade",
                   "Resumo"]
        widths  = [32, 22, 12, 12, 8, 10, 10, 10, 14, 10, 12, 12, 12, 12, 60]

    elif area_key == "COMERCIAL_CADASTRO":
        headers = ["ID", "Agente", "Data", "Duração (s)", "Efetivo", "Coaching",
                   "Tipo Contato", "Resultado", "Próximo Passo",
                   "Impacto Receita", "Conformidade", "Resumo"]
        widths  = [32, 22, 12, 12, 10, 10, 18, 16, 28, 15, 12, 60]
    else:
        headers = ["ID", "Agente", "Data", "Duração (s)", "Resumo"]
        widths  = [32, 22, 12, 12, 80]

    # Cabeçalho
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        _style_header(cell)
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 30

    # Dados
    for i, lig in enumerate(ligacoes, 2):
        row_bg = COR_LINHA_PAR if i % 2 == 0 else "FFFFFF"

        if area_key == "SUCESSO_CORRETOR":
            spin = lig.get("spin_aplicado", {})
            vals = [
                lig.get("id"), lig.get("agente"), lig.get("data"),
                lig.get("duracao_s"), lig.get("efetivo"), lig.get("prioridade_coaching"),
                lig.get("tipo_pendencia_detectado", "-"),
                lig.get("padrao_comercial", "-"),
                lig.get("impacto_receita", "-"),
                spin.get("situacao", "-"), spin.get("problema", "-"),
                spin.get("implicacao", "-"), spin.get("necessidade", "-"),
                lig.get("conformidade", "-"), lig.get("resumo", "-"),
            ]
            cor_efet = _cor_efetividade(lig.get("efetivo"))
            cor_coach = _cor_coaching(lig.get("prioridade_coaching"))

        elif area_key == "ATENDIMENTO":
            crit = lig.get("criterios", {})
            def c(k): return "1" if crit.get(k, {}).get("status") == "CUMPRIU" else "-"
            vals = [
                lig.get("id"), lig.get("agente"), lig.get("data"),
                lig.get("duracao_s"), lig.get("nota_final", "-"),
                lig.get("prioridade_coaching"),
                c("pronto_atendimento_saudacao"), c("atencao_concentracao"),
                c("entrega_solucao_efetiva"), c("clareza_seguranca"),
                c("transferencia_contato"), c("tempo_espera_resposta"),
                c("vocabulario"), c("cordialidade_empatia"),
                lig.get("resumo", "-"),
            ]
            cor_efet = None
            cor_coach = _cor_coaching(lig.get("prioridade_coaching"))

        elif area_key == "COMERCIAL_CADASTRO":
            res = lig.get("resultado_comercial", {})
            vals = [
                lig.get("id"), lig.get("agente"), lig.get("data"),
                lig.get("duracao_s"), lig.get("efetivo"), lig.get("prioridade_coaching"),
                lig.get("tipo_contato", "-"),
                res.get("status", "-"), res.get("proximo_passo", "-"),
                lig.get("impacto_receita", "-"),
                lig.get("conformidade", "-"), lig.get("resumo", "-"),
            ]
            cor_efet = _cor_efetividade(lig.get("efetivo"))
            cor_coach = _cor_coaching(lig.get("prioridade_coaching"))
        else:
            vals = [lig.get("id"), lig.get("agente"), lig.get("data"),
                    lig.get("duracao_s"), lig.get("resumo", "-")]
            cor_efet = cor_coach = None

        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.font = Font(name="Arial", size=9)
            cell.alignment = Alignment(vertical="center", wrap_text=(col == len(vals)))
            cell.fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")
            _bordas(cell)

        # Cor condicional em Efetivo e Coaching
        col_efet = headers.index("Efetivo") + 1 if "Efetivo" in headers else None
        col_coach = headers.index("Coaching") + 1 if "Coaching" in headers else None
        if cor_efet and col_efet:
            ws.cell(row=i, column=col_efet).fill = PatternFill(start_color=cor_efet, end_color=cor_efet, fill_type="solid")
        if cor_coach and col_coach:
            ws.cell(row=i, column=col_coach).fill = PatternFill(start_color=cor_coach, end_color=cor_coach, fill_type="solid")

        ws.row_dimensions[i].height = 22

    return ws


# ── ABA POR GESTOR ─────────────────────────────────────────────────────────────

def aba_por_gestor(wb, relatorio):
    ws = wb.create_sheet("Por Gestor")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    area_key = relatorio.get("area_key", "")
    gestores = relatorio.get("ranking_gestores", [])

    if area_key == "SUCESSO_CORRETOR":
        headers = ["Gestor", "Total", "Efetivas", "Efetividade %", "Coaching Predominante", "SPIN Médio"]
        widths  = [28, 10, 10, 14, 22, 12]
        def vals(g): return [
            g.get("agente"), g.get("total"), g.get("efetivas"),
            f"{g.get('efetividade_pct', 0):.1f}%",
            g.get("coaching_predominante", "-"), g.get("spin_medio", "-"),
        ]
    elif area_key == "ATENDIMENTO":
        headers = ["Gestor", "Total", "Nota Média", "Coaching Predominante", "Critério Mais Falho"]
        widths  = [28, 10, 12, 22, 40]
        def vals(g): return [
            g.get("agente"), g.get("total"),
            g.get("nota_media", "-"),
            g.get("coaching_predominante", "-"),
            g.get("criterio_mais_falho", "-"),
        ]
    elif area_key == "COMERCIAL_CADASTRO":
        headers = ["Gestor", "Total", "Convertidas", "Conversão %", "Coaching Predominante", "SPIN Médio"]
        widths  = [28, 10, 12, 12, 22, 12]
        def vals(g): return [
            g.get("agente"), g.get("total"), g.get("convertidas"),
            f"{g.get('conversao_pct', 0):.1f}%",
            g.get("coaching_predominante", "-"), g.get("spin_medio", "-"),
        ]
    else:
        return ws

    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col, value=h)
        _style_header(cell)
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.row_dimensions[1].height = 28

    for i, g in enumerate(gestores, 2):
        row_bg = COR_LINHA_PAR if i % 2 == 0 else "FFFFFF"
        for col, val in enumerate(vals(g), 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="center", horizontal="center")
            cell.fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")
            _bordas(cell)
        ws.row_dimensions[i].height = 20

    return ws


# ── ABA FALHAS COMUNS ──────────────────────────────────────────────────────────

def aba_falhas(wb, relatorio):
    ws = wb.create_sheet("Falhas Comuns")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 14

    ws.merge_cells("A1:B1")
    ws["A1"].value = "FALHAS MAIS FREQUENTES"
    ws["A1"].font = Font(bold=True, size=11, color="FFFFFF", name="Arial")
    ws["A1"].fill = PatternFill(start_color=COR_VERMELHO.replace("FCE4D6", "C00000"), end_color="C00000", fill_type="solid")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    for col, h in enumerate(["Problema Identificado", "Ocorrências"], 1):
        cell = ws.cell(row=2, column=col, value=h)
        _style_header(cell, cor=COR_SUBHEADER)
    ws.row_dimensions[2].height = 22

    for i, falha in enumerate(relatorio.get("falhas_comuns", []), 3):
        row_bg = COR_LINHA_PAR if i % 2 == 0 else "FFFFFF"
        ws.cell(row=i, column=1, value=falha.get("item", "")).font = Font(name="Arial", size=10)
        ws.cell(row=i, column=2, value=falha.get("ocorrencias", 0)).font = Font(bold=True, name="Arial", size=10)
        ws.cell(row=i, column=1).fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")
        ws.cell(row=i, column=2).fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")
        ws.cell(row=i, column=2).alignment = Alignment(horizontal="center")
        for col in [1, 2]:
            _bordas(ws.cell(row=i, column=col))
        ws.row_dimensions[i].height = 18

    # Boas práticas
    row_base = len(relatorio.get("falhas_comuns", [])) + 5
    ws.merge_cells(f"A{row_base}:B{row_base}")
    ws[f"A{row_base}"].value = "BOAS PRÁTICAS PARA REPLICAR"
    ws[f"A{row_base}"].font = Font(bold=True, size=11, color="FFFFFF", name="Arial")
    ws[f"A{row_base}"].fill = PatternFill(start_color="375623", end_color="375623", fill_type="solid")
    ws[f"A{row_base}"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row_base].height = 24

    row_base += 1
    for col, h in enumerate(["Prática Identificada", "Ocorrências"], 1):
        cell = ws.cell(row=row_base, column=col, value=h)
        _style_header(cell, cor="548235")
    ws.row_dimensions[row_base].height = 22

    for i, pratica in enumerate(relatorio.get("boas_praticas_replicaveis", []), row_base + 1):
        row_bg = COR_VERDE if i % 2 == 0 else "FFFFFF"
        ws.cell(row=i, column=1, value=pratica.get("item", "")).font = Font(name="Arial", size=10)
        ws.cell(row=i, column=2, value=pratica.get("ocorrencias", 0)).font = Font(bold=True, name="Arial", size=10)
        for col in [1, 2]:
            ws.cell(row=i, column=col).fill = PatternFill(start_color=row_bg, end_color=row_bg, fill_type="solid")
            ws.cell(row=i, column=col).alignment = Alignment(horizontal="center" if col == 2 else "left")
            _bordas(ws.cell(row=i, column=col))
        ws.row_dimensions[i].height = 18

    return ws


# ── MAIN ───────────────────────────────────────────────────────────────────────

def gerar(data=None):
    if not data:
        # Pegar pasta mais recente
        pastas = sorted(Path(OUTPUT_DIR).iterdir(), reverse=True)
        pasta = next((p for p in pastas if p.is_dir()), None)
    else:
        pasta = Path(OUTPUT_DIR) / data

    if not pasta or not pasta.exists():
        print(f"Pasta de output não encontrada: {pasta}")
        return

    jsons = sorted(pasta.glob("relatorio_*.json"))
    if not jsons:
        print(f"Nenhum relatório JSON encontrado em {pasta}")
        return

    print(f"Gerando Excel a partir de {len(jsons)} relatório(s) em {pasta.name}...")

    for json_file in jsons:
        with open(json_file, encoding="utf-8") as f:
            relatorio = json.load(f)

        area_key = relatorio.get("area_key", "DESCONHECIDO")
        wb = Workbook()
        wb.remove(wb.active)

        aba_resumo(wb, relatorio)
        aba_por_ligacao(wb, relatorio)
        aba_por_gestor(wb, relatorio)
        aba_falhas(wb, relatorio)

        nome_excel = json_file.name.replace(".json", ".xlsx")
        caminho = pasta / nome_excel
        wb.save(str(caminho))
        print(f"  Salvo: {caminho}")

    print("\nConcluido!")


if __name__ == "__main__":
    data = sys.argv[1] if len(sys.argv) > 1 else None
    gerar(data)
